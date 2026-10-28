"""Step 4: export the results to a CSV (opens in Excel / Google Sheets) and a
Word file that shows each scan next to its transcription.

Usage:
    python step4_export.py questions_2560
    python step4_export.py questions_2560 --approved-only

Files are written to work/<pdf name>/export/.
"""

import argparse
import csv

from docx import Document
from docx.shared import Cm, Pt

from common import best_version, list_pages, work_dir_for

CSV_COLUMNS = ["page", "status", "confidence", "greeting", "question", "closing",
               "lines", "spelling_fixes", "notes", "continues_from_previous", "continues_to_next",
               "reviewed_by", "reviewed_at"]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("document")
    parser.add_argument("--approved-only", action="store_true", help="leave out pages not approved yet")
    args = parser.parse_args()

    doc_dir = work_dir_for(args.document)
    export_dir = doc_dir / "export"
    export_dir.mkdir(exist_ok=True)

    rows = []
    for page in list_pages(doc_dir):
        data, status = best_version(doc_dir, page)
        if data is None or (args.approved_only and status != "approved"):
            continue
        rows.append({**data, "page": int(page.split("_")[1]), "status": status,
                     "lines": "\n".join(data.get("lines", [])),
                     "spelling_fixes": "; ".join(f"{fix['as_written']} → {fix['corrected']}"
                                                 for fix in data.get("corrections", []))})

    # utf-8-sig adds a marker so Excel shows Thai correctly instead of garbled letters.
    csv_path = export_dir / f"{args.document}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    word = Document()
    word.styles["Normal"].font.size = Pt(14)
    word.add_heading(args.document, level=0)
    for row in rows:
        page = f"page_{row['page']:03d}"
        word.add_heading(f"Page {row['page']}  ({row['status']})", level=1)
        table = word.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        image_cell, text_cell = table.rows[0].cells
        image_cell.paragraphs[0].add_run().add_picture(str(doc_dir / "crops" / f"{page}.jpg"), width=Cm(8))
        text_cell.paragraphs[0].add_run(row["lines"])
        if row.get("question"):
            word.add_paragraph().add_run("Clean version: ").bold = True
            word.add_paragraph(" ".join(filter(None, [row.get("greeting"), row["question"], row.get("closing")])))
        if row.get("spelling_fixes"):
            word.add_paragraph(f"Spelling fixed (as written → corrected): {row['spelling_fixes']}")
        if row.get("notes"):
            word.add_paragraph(f"Notes: {row['notes']}")
        word.add_page_break()
    docx_path = export_dir / f"{args.document}.docx"
    word.save(docx_path)

    print(f"Exported {len(rows)} page(s):\n  {csv_path}\n  {docx_path}")


if __name__ == "__main__":
    main()
