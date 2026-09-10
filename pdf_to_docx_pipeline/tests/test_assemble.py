import io

import cv2
import numpy as np
from docx import Document

from pipeline import assemble


def _tiny_png_bytes():
    img = np.full((20, 20, 3), 255, dtype=np.uint8)
    ok, buf = cv2.imencode(".png", img)
    assert ok
    return buf.tobytes()


def test_build_docx_writes_text_image_and_handwriting_items(tmp_path):
    items = [
        {"kind": "heading", "page": 0},
        {"kind": "text", "lang": "en", "text": "Hello world"},
        {"kind": "image", "image_bytes": _tiny_png_bytes()},
        {"kind": "handwriting", "image_bytes": _tiny_png_bytes()},
        {"kind": "uncertain", "image_bytes": _tiny_png_bytes()},
    ]
    out = tmp_path / "out.docx"
    assemble.build_docx(items, str(out))

    assert out.exists()
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Hello world" in text
    assert "[EN]" in text
    assert "Handwritten" in text
    assert "Low-confidence OCR" in text
    assert len(doc.inline_shapes) == 3


def test_build_docx_strips_control_characters(tmp_path):
    items = [{"kind": "text", "lang": "en", "text": "Bad\x00Char\x07Here"}]
    out = tmp_path / "out.docx"
    assemble.build_docx(items, str(out))
    doc = Document(str(out))
    joined = "".join(p.text for p in doc.paragraphs)
    assert "\x00" not in joined
    assert "\x07" not in joined
    assert "BadCharHere" in joined


def test_crop_to_png_bytes_respects_padding_and_bounds():
    img = np.zeros((50, 50), dtype=np.uint8)
    cropped = assemble.crop_to_png_bytes(img, (10, 10, 20, 20), pad=6)
    decoded = cv2.imdecode(np.frombuffer(cropped, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
    assert decoded.shape == (22, 22)  # (20-10)+2*6 in each dim, clipped to image bounds


def test_crop_to_png_bytes_clips_at_image_edges():
    img = np.zeros((30, 30), dtype=np.uint8)
    cropped = assemble.crop_to_png_bytes(img, (0, 0, 5, 5), pad=10)
    decoded = cv2.imdecode(np.frombuffer(cropped, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
    assert decoded.shape == (15, 15)  # top/left clipped to 0, bottom/right padded normally
